"""Document and data conversion utility supporting DOCX, XLSX, PDF, CSV, TXT, HTML, JSON."""

import os
import sys
import subprocess
from typing import Callable, Dict, Optional, Tuple, List
try:
    import pandas as pd
except ImportError:
    pd = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError:
    pass

# Compatibility patch for pdf2docx with PyMuPDF >= 1.25
try:
    import fitz
    if hasattr(fitz, "Rect") and not hasattr(fitz.Rect, "get_area"):
        fitz.Rect.get_area = lambda self: abs(self)
except Exception:
    pass


DOCUMENT_CONVERSIONS = {
    # Source extensions mapped to allowable target formats
    "pdf": ["DOCX", "TXT", "PNG", "HTML"],
    "docx": ["PDF", "TXT", "HTML", "RTF"],
    "xlsx": ["CSV", "PDF", "JSON", "HTML"],
    "xls": ["CSV", "XLSX", "PDF"],
    "csv": ["XLSX", "PDF", "JSON", "HTML", "TXT"],
    "txt": ["PDF", "DOCX", "HTML"],
    "html": ["PDF", "DOCX", "TXT"],
    "rtf": ["DOCX", "PDF", "TXT", "HTML"],
}


def get_available_targets(source_path: str) -> List[str]:
    """Return available target formats for a given source document."""
    ext = os.path.splitext(source_path)[1].lower().lstrip(".")
    return DOCUMENT_CONVERSIONS.get(ext, ["PDF", "TXT"])


def convert_pdf_to_docx(input_file: str, output_file: str, progress_cb: Optional[Callable[[float, str], None]] = None) -> Tuple[bool, str]:
    """Convert PDF to editable DOCX with layout preservation."""
    try:
        from pdf2docx import Converter
        if progress_cb:
            progress_cb(20.0, "Initialisiere PDF-Konvertierung...")
        cv = Converter(input_file)
        if progress_cb:
            progress_cb(50.0, "Konvertiere Seiten und Tabellen...")
        cv.convert(output_file, start=0, end=None)
        cv.close()
        return True, "PDF erfolgreich in DOCX umgewandelt."
    except Exception as e:
        return False, f"Fehler bei PDF-zu-DOCX: {str(e)}"


def convert_pdf_to_txt(input_file: str, output_file: str, progress_cb: Optional[Callable[[float, str], None]] = None) -> Tuple[bool, str]:
    """Extract plain text from PDF."""
    try:
        if progress_cb:
            progress_cb(25.0, "Lese PDF-Text...")
        reader = PdfReader(input_file)
        text_parts = []
        total = len(reader.pages)
        for i, page in enumerate(reader.pages):
            text_parts.append(page.extract_text() or "")
            if progress_cb and total > 0:
                progress_cb(min(90.0, 25.0 + (i / total) * 65.0), f"Seite {i+1} von {total}...")
        
        full_text = "\n\n--- Seite getrennt ---\n\n".join(text_parts)
        with open(output_file, "w", encoding="utf-8") as f:
            f.write(full_text)
        return True, "PDF-Text erfolgreich exportiert."
    except Exception as e:
        return False, f"Fehler beim PDF-Text-Export: {str(e)}"


def convert_pdf_to_images(input_file: str, output_file: str, progress_cb: Optional[Callable[[float, str], None]] = None) -> Tuple[bool, str]:
    """Convert PDF pages to PNG image files."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(input_file)
        total = len(doc)
        base, _ = os.path.splitext(output_file)
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=200)
            target_path = output_file if total == 1 else f"{base}_seite_{i+1}.png"
            pix.save(target_path)
            if progress_cb and total > 0:
                progress_cb(min(90.0, (i + 1) / total * 90.0), f"Rendere Seite {i+1}/{total}...")
        return True, f"{total} PDF-Seiten als Bild gerendert."
    except Exception as e:
        return False, f"Fehler beim Rendern von PDF: {str(e)}"


def convert_table_file(input_file: str, output_file: str, target_fmt: str, progress_cb: Optional[Callable[[float, str], None]] = None) -> Tuple[bool, str]:
    """Convert tabular files (CSV, XLSX, XLS) between formats."""
    src_ext = os.path.splitext(input_file)[1].lower()
    target_fmt = target_fmt.upper()

    try:
        if progress_cb:
            progress_cb(20.0, "Lese Tabelle...")

        if src_ext in [".xlsx", ".xls"]:
            df = pd.read_excel(input_file)
        elif src_ext in [".csv", ".tsv"]:
            # Auto-detect separator
            try:
                df = pd.read_csv(input_file, sep=None, engine='python')
            except Exception:
                df = pd.read_csv(input_file)
        else:
            return False, f"Nicht unterstütztes Quellformat: {src_ext}"

        if progress_cb:
            progress_cb(60.0, f"Erzeuge {target_fmt}...")

        if target_fmt == "CSV":
            df.to_csv(output_file, index=False, encoding="utf-8-sig")
        elif target_fmt == "XLSX":
            df.to_excel(output_file, index=False, engine="openpyxl")
        elif target_fmt == "JSON":
            df.to_json(output_file, orient="records", indent=2, force_ascii=False)
        elif target_fmt == "HTML":
            df.to_html(output_file, index=False)
        elif target_fmt == "PDF":
            # Generate table PDF with reportlab
            doc = SimpleDocTemplate(output_file, pagesize=A4, rightMargin=20, leftMargin=20, topMargin=20, bottomMargin=20)
            elements = []
            
            # Limit width/cols for clean rendering
            preview_df = df.head(100) # up to 100 rows
            data = [[str(col)[:25] for col in preview_df.columns]]
            for _, row in preview_df.iterrows():
                data.append([str(val)[:30] for val in row.values])

            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ]))
            elements.append(t)
            doc.build(elements)

        if progress_cb:
            progress_cb(100.0, "Fertig!")
        return True, f"Erfolgreich in {target_fmt} konvertiert."
    except Exception as e:
        return False, f"Tabellen-Konvertierungsfehler: {str(e)}"


def convert_docx_or_text(input_file: str, output_file: str, target_fmt: str, progress_cb: Optional[Callable[[float, str], None]] = None) -> Tuple[bool, str]:
    """Convert text documents (DOCX, TXT, RTF, HTML)."""
    src_ext = os.path.splitext(input_file)[1].lower().lstrip(".")
    target_fmt = target_fmt.upper()

    try:
        if progress_cb:
            progress_cb(20.0, "Verarbeite Dokument...")

        # Case 1: To PDF using ReportLab
        if target_fmt == "PDF":
            text_content = ""
            if src_ext == "docx":
                import docx
                doc = docx.Document(input_file)
                text_content = "\n\n".join([p.text for p in doc.paragraphs if p.text])
            else:
                with open(input_file, "r", encoding="utf-8", errors="ignore") as f:
                    text_content = f.read()

            pdf_doc = SimpleDocTemplate(output_file, pagesize=A4)
            styles = getSampleStyleSheet()
            normal_style = styles["Normal"]
            elements = []
            for line in text_content.splitlines():
                if line.strip():
                    elements.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), normal_style))
                    elements.append(Spacer(1, 4))
            pdf_doc.build(elements)
            return True, "Dokument erfolgreich als PDF gespeichert."

        # Native Python fallback for TXT/MD -> DOCX
        if target_fmt == "DOCX" and src_ext in ("txt", "md"):
            import docx
            doc = docx.Document()
            with open(input_file, "r", encoding="utf-8", errors="ignore") as f:
                for line in f:
                    doc.add_paragraph(line.rstrip("\r\n"))
            doc.save(output_file)
            return True, "Erfolgreich in DOCX konvertiert."

        # Case 2: Using macOS textutil (native between docx, html, txt, rtf on macOS)
        if sys.platform == "darwin":
            target_mapping = {"DOCX": "docx", "TXT": "txt", "HTML": "html", "RTF": "rtf"}
            if target_fmt in target_mapping:
                out_fmt = target_mapping[target_fmt]
                cmd = ["textutil", "-convert", out_fmt, "-output", output_file, input_file]
                res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                if res.returncode == 0 and os.path.exists(output_file):
                    return True, f"Erfolgreich in {target_fmt} konvertiert."

        # Fallback for plain text
        if target_fmt == "TXT":
            if src_ext == "docx":
                import docx
                doc = docx.Document(input_file)
                content = "\n".join([p.text for p in doc.paragraphs])
                with open(output_file, "w", encoding="utf-8") as f:
                    f.write(content)
                return True, "DOCX erfolgreich in TXT umgewandelt."

        return False, f"Konvertierung von {src_ext} nach {target_fmt} wird nicht unterstützt."
    except Exception as e:
        return False, f"Fehler bei Dokumenten-Konvertierung: {str(e)}"


def convert_document(
    input_file: str,
    output_file: str,
    target_format: str,
    progress_callback: Optional[Callable[[float, str], None]] = None,
) -> Tuple[bool, str]:
    """Master entry point for document and spreadsheet conversions."""
    if not os.path.isfile(input_file):
        return False, f"Datei existiert nicht: {input_file}"

    src_ext = os.path.splitext(input_file)[1].lower().lstrip(".")
    target_fmt = target_format.upper()

    out_dir = os.path.dirname(output_file)
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    # Route based on source and target
    if src_ext == "pdf":
        if target_fmt == "DOCX":
            return convert_pdf_to_docx(input_file, output_file, progress_callback)
        elif target_fmt == "TXT":
            return convert_pdf_to_txt(input_file, output_file, progress_callback)
        elif target_fmt in ["PNG", "JPG", "JPEG"]:
            return convert_pdf_to_images(input_file, output_file, progress_callback)

    if src_ext in ["xlsx", "xls", "csv", "tsv"]:
        return convert_table_file(input_file, output_file, target_fmt, progress_callback)

    if src_ext in ["docx", "txt", "html", "rtf"]:
        return convert_docx_or_text(input_file, output_file, target_fmt, progress_callback)

    return False, f"Keine passende Konvertierung für .{src_ext} nach {target_fmt} gefunden."
